from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(r"C:\Users\Lenovo\Desktop\the-old-no-deps")
SOURCE_DOC = Path(r"C:\Users\Lenovo\Desktop\提交文档\1毕业论文\2022213086-汪序磊-毕业设计降重版.docx")
OUTPUT_DOC = SOURCE_DOC.with_name("2022213086-汪序磊-毕业设计降重版-第五章代码展示.docx")


def read_lines(relative_path: str, start: int, end: int) -> str:
    path = ROOT / relative_path
    lines = path.read_text(encoding="utf-8").splitlines()
    return "\n".join(lines[start - 1 : end])


def set_paragraph_font(paragraph, font_name: str, size_pt: float, bold: bool = False) -> None:
    for run in paragraph.runs:
        run.font.name = font_name
        run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
        run.font.size = Pt(size_pt)
        run.bold = bold


def insert_paragraph_after(paragraph, text: str):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = paragraph._parent.add_paragraph()
    new_para._p.getparent().remove(new_para._p)
    new_p.addnext(new_para._p)
    new_para._p.getparent().remove(new_p)
    new_para.add_run(text)
    return new_para


SECTION_CODE = {
    "5.1": {
        "intro": "其中实现JWT身份验证逻辑的代码如下：",
        "code": "\n\n".join(
            [
                "const token = jwt.sign(",
                "  { user: { id: userData.id, role: userData.role, status: userData.status } },",
                "  JWT_SECRET,",
                "  { expiresIn: '7d' }",
                ");",
                "",
                "const authHeader = req.header('Authorization');",
                "if (!authHeader || !authHeader.startsWith('Bearer ')) {",
                "  return res.status(401).json({ message: 'No token, authorization denied' });",
                "}",
                "const token = authHeader.split(' ')[1];",
                "const decoded = jwt.verify(token, JWT_SECRET);",
                "req.user = decoded.user;",
            ]
        ),
    },
    "5.2": {
        "intro": "其中实现首页数据加载与状态统计的关键代码如下：",
        "code": "\n\n".join(
            [
                "const fetchOrders = async () => {",
                "  setLoading(true);",
                "  const res = await api.get('/orders');",
                "  const sortedOrders = res.data.orders.sort((a, b) =>",
                "    new Date(b.createdAt).getTime() - new Date(a.createdAt).getTime()",
                "  );",
                "  setOrders(sortedOrders);",
                "};",
                "",
                "const inProgressCount = orders.filter(o => ['accepted', 'in_progress'].includes(o.status)).length;",
                "const pendingCount = orders.filter(o => ['under_review', 'pending'].includes(o.status)).length;",
                "const completedCount = orders.filter(o => o.status === 'completed').length;",
            ]
        ),
    },
    "5.3": {
        "intro": "其中实现需求发布与订单创建逻辑的关键代码如下：",
        "code": "\n\n".join(
            [
                "if (values.expectedTime.isBefore(dayjs())) {",
                "  Modal.warning({ title: '上门时间填写错误' });",
                "  return;",
                "}",
                "await api.post('/orders', {",
                "  ...values,",
                "  expectedTime: values.expectedTime.toISOString(),",
                "  lat: selectedLocation.lat,",
                "  lng: selectedLocation.lng",
                "});",
                "",
                "const newOrder = await Order.create({",
                "  title,",
                "  category,",
                "  description,",
                "  address,",
                "  expectedTime,",
                "  lat,",
                "  lng,",
                "  elderlyId: userId,",
                "  status: 'under_review'",
                "});",
            ]
        ),
    },
    "5.4": {
        "intro": "其中实现任务地图加载、附近任务筛选与接单逻辑的关键代码如下：",
        "code": "\n\n".join(
            [
                "const fetchOrders = async (lat?: number, lng?: number) => {",
                "  const params = { tab: 'available' };",
                "  if (lat !== undefined && lng !== undefined) {",
                "    params.lat = lat;",
                "    params.lng = lng;",
                "    params.radius = 15;",
                "  }",
                "  const { data } = await api.get('/orders', { params });",
                "  setOrders(data.orders);",
                "};",
                "",
                "const handleAcceptOrder = async (orderId: number) => {",
                "  await api.post(`/orders/${orderId}/accept`);",
                "  message.success('接单成功，快去服务吧');",
                "  navigate('/volunteer');",
                "};",
            ]
        ),
    },
    "5.5": {
        "intro": "其中实现服务状态更新与积分结算逻辑的关键代码如下：",
        "code": "\n\n".join(
            [
                "order.status = status;",
                "if (completionPhotos) {",
                "  order.completionPhotos = JSON.stringify(completionPhotos);",
                "}",
                "if (completionDescription) {",
                "  order.completionDescription = completionDescription;",
                "}",
                "await order.save();",
                "",
                "const pointsToAward = pointsReward != null ? pointsReward : (POINT_RULES[category] || POINT_RULES['default']);",
                "const user = await User.findByPk(volunteerId, { transaction });",
                "user.points += pointsToAward;",
                "await user.save({ transaction });",
                "await PointLog.create({",
                "  userId: volunteerId,",
                "  pointsChanged: pointsToAward,",
                "  reason: `Completed order: ${category}`,",
                "  orderId",
                "}, { transaction });",
            ]
        ),
    },
    "5.6": {
        "intro": "其中实现管理员数据看板聚合统计接口的关键代码如下：",
        "code": "\n\n".join(
            [
                "const pendingOrdersCount = await Order.count({ where: { status: 'pending' } });",
                "const totalMonthOrders = await Order.count({ where: { createdAt: { [Op.gte]: monthStart } } });",
                "const completedMonthOrders = await Order.count({",
                "  where: {",
                "    createdAt: { [Op.gte]: monthStart },",
                "    status: 'completed'",
                "  }",
                "});",
                "const completionRate = totalMonthOrders === 0 ? 0 : (completedMonthOrders / totalMonthOrders) * 100;",
            ]
        ),
    },
    "5.7": {
        "intro": "其中实现志愿者审核与账号状态维护的关键代码如下：",
        "code": "\n\n".join(
            [
                "const handleStatusUpdate = async (id: number, status: string) => {",
                "  await api.put(`/admin/users/${id}/status`, { status });",
                "  message.success('状态更新成功');",
                "  await fetchUsers();",
                "};",
                "",
                "const user = await User.findByPk(parseInt(id as string, 10));",
                "if (!user) {",
                "  return res.status(404).json({ message: 'User not found' });",
                "}",
                "user.status = status;",
                "await user.save();",
            ]
        ),
    },
    "5.8": {
        "intro": "其中实现需求审核、赋分与工单调度的关键代码如下：",
        "code": "\n\n".join(
            [
                "await api.put(`/admin/orders/${selectedOrder.id}/audit`, {",
                "  status: 'pending',",
                "  pointsReward,",
                "});",
                "",
                "if (status !== 'pending' && status !== 'rejected' && status !== 'submitted') {",
                "  return res.status(400).json({ message: 'Invalid audit status' });",
                "}",
                "if (!isCompletionDispute && status === 'pending' && (pointsReward === undefined || pointsReward === null || pointsReward <= 0)) {",
                "  return res.status(400).json({ message: 'Approved orders must have a valid points reward' });",
                "}",
                "order.status = status;",
                "if (status === 'pending') {",
                "  order.pointsReward = pointsReward;",
                "}",
                "await order.save();",
            ]
        ),
    },
    "5.9": {
        "intro": "其中实现统计分析图表渲染的关键代码如下：",
        "code": "\n\n".join(
            [
                "const data = stats.categoriesData.map((d) => ({",
                "  name: d.category,",
                "  value: d.count",
                "}));",
                "",
                "return {",
                "  title: { text: '需求类型分布', left: 'center' },",
                "  tooltip: { trigger: 'item' },",
                "  legend: { orient: 'vertical', left: 'left' },",
                "  series: [{",
                "    name: '需求分类',",
                "    type: 'pie',",
                "    radius: '50%',",
                "    data,",
                "  }]",
                "};",
            ]
        ),
    },
    "5.10": {
        "intro": "其中实现个人信息与老人档案读取、更新的关键代码如下：",
        "code": "\n\n".join(
            [
                "const user = await User.findByPk(req.user.id, {",
                "  attributes: ['id', 'name', 'phone', 'address', 'role', 'points'],",
                "});",
                "const redemptions = await PrizeRedemption.findAll({",
                "  where: { targetPhone: user.phone },",
                "  order: [['createdAt', 'DESC']]",
                "});",
                "userData.prizeRedemptions = redemptions;",
                "",
                "user.name = name !== undefined ? name : user.name;",
                "user.phone = phone !== undefined ? phone : user.phone;",
                "user.address = address !== undefined ? address : user.address;",
                "await user.save();",
            ]
        ),
    },
}


def find_section_anchors(document: Document):
    paragraphs = list(document.paragraphs)
    anchors = {}
    for idx, paragraph in enumerate(paragraphs):
        heading = paragraph.text.strip()
        if not heading.startswith("5."):
            continue

        section = heading.split()[0]
        if section not in SECTION_CODE:
            continue

        end_idx = len(paragraphs)
        for j in range(idx + 1, len(paragraphs)):
            next_text = paragraphs[j].text.strip()
            if next_text.startswith("5.") or next_text.startswith("第6章"):
                end_idx = j
                break

        for candidate in reversed(paragraphs[idx + 1 : end_idx]):
            if candidate.text.strip():
                anchors[section] = candidate
                break
    return anchors


def main() -> None:
    shutil.copy2(SOURCE_DOC, OUTPUT_DOC)
    doc = Document(OUTPUT_DOC)
    anchors = find_section_anchors(doc)
    missing = sorted(set(SECTION_CODE) - set(anchors))
    if missing:
        raise RuntimeError(f"未找到以下章节插入点: {', '.join(missing)}")

    for section in sorted(SECTION_CODE.keys(), key=lambda value: [int(part) for part in value.split(".")], reverse=True):
        item = SECTION_CODE[section]
        anchor = anchors[section]

        intro = insert_paragraph_after(anchor, item["intro"])
        intro.paragraph_format.first_line_indent = Pt(21)
        intro.paragraph_format.space_before = Pt(6)
        intro.paragraph_format.space_after = Pt(3)
        set_paragraph_font(intro, "宋体", 10.5)

        code = insert_paragraph_after(intro, item["code"])
        code.paragraph_format.first_line_indent = Pt(0)
        code.paragraph_format.space_before = Pt(0)
        code.paragraph_format.space_after = Pt(6)
        code.paragraph_format.line_spacing = 1.0
        set_paragraph_font(code, "Consolas", 9)

    doc.save(OUTPUT_DOC)
    print(OUTPUT_DOC)


if __name__ == "__main__":
    main()
